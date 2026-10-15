"""docs/SPEC-v4.md 4.4 acceptance tests D1-D5: staff run daily orders safely; only the owner
changes money settings and makes learning permanent."""
import io
import re

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.db import SessionLocal
from app.models import User
from tests.test_cart import deliver, say, shop, tap  # noqa: F401  (fixture)
from tests.test_messaging import scalar
from tests.test_payments import downloaded, order, photo, waiting_for_payment
from tests.test_v4_teaching import learned, pid, staff_client, teach, y_shop
from tests.test_web import PASSWORD, client, csrf, log_in  # noqa: F401  (fixture)

OWNER_GETS = ("/prices", "/prices/export.xlsx", "/settings", "/learning", "/users", "/activity")
OWNER_POSTS = (
    "/prices/import", "/prices/import/confirm", "/settings/payment-methods", "/settings/payment-methods/1",
    "/settings/payment-methods/1/delete", "/settings/templates/greeting", "/settings/templates/greeting/reset",
    "/settings/faq", "/settings/faq/1", "/settings/faq/1/delete", "/settings/order-rules", "/settings/vial-costs",
    "/settings/auto-speed", "/settings/learning", "/auto-reply", "/users", "/users/1/password", "/users/1/disabled",
    "/users/1/language", "/users/staff-settings", "/learning/1/decide", "/learning/1/edit", "/learning/conflicts/1",
    "/learning/price/1", "/products", "/products/1/active", "/products/1/aliases", "/products/1/aliases/1/delete",
    "/sales/fill-costs", "/messages/1/unknown/0/map", "/messages/1/unknown/0/new",
)
OWNER_LINKS = ('href="/prices"', 'href="/settings"', 'href="/learning"', 'href="/users"', 'action="/auto-reply"')


def test_d1_staff_gets_403_and_no_menu_links(shop, db):
    owner, _ = shop
    staff, token = staff_client(db)
    for path in OWNER_GETS:
        r = staff.get(path)
        assert r.status_code == 403 and "Only the owner" in r.text, path
        assert owner.get(path).status_code == 200, path
    for path in OWNER_POSTS:
        r = staff.post(path, data={"csrf": token}, follow_redirects=False)
        assert r.status_code == 403 and "Only the owner" in r.text, path  # not the CSRF 403
    home = staff.get("/").text
    assert not any(link in home for link in OWNER_LINKS)
    assert "🤖 Auto-reply: <b>" in home  # staff still see whether auto-reply is on
    owner_home = owner.get("/").text
    assert all(link in owner_home for link in OWNER_LINKS)
    # what staff CAN do: inbox, chats, orders, payments list, reports, products (read only), sales
    for path in ("/", "/chats", "/orders", "/payments", "/reports/wanted", "/products", "/sales", "/health"):
        assert staff.get(path).status_code == 200, path
    assert 'method="post" action="/products"' not in staff.get("/products").text


def test_d1b_staff_sales_without_cost_and_profit(shop, db):
    owner, _ = shop
    staff, _ = staff_client(db)
    page = staff.get("/sales").text
    assert "Revenue" in page and "Orders" in page and "Profit" not in page and "Cost" not in page
    assert "Profit" in owner.get("/sales").text
    wb = load_workbook(io.BytesIO(staff.get("/sales/export.xlsx").content))
    summary = [row[0] for row in wb["Summary"].iter_rows(values_only=True)]
    assert "Revenue (ကျပ်)" in summary and not any(str(k).startswith(("Cost", "Profit")) for k in summary)
    assert list(wb["Per day"].iter_rows(values_only=True))[0] == ("Day", "Orders", "Items", "ml", "Revenue")


def _photo_check(owner, owner_token, db, customer: str, n: int) -> int:
    waiting_for_payment(owner, owner_token, db, customer=customer)
    photo(n, customer=customer)
    return scalar(db, "SELECT max(id) FROM payment_checks")


def test_d2_staff_payments_follow_the_setting(shop, db, tmp_path, monkeypatch):
    owner, owner_token = shop
    staff, token = staff_client(db)
    check = _photo_check(owner, owner_token, db, "91", 1)
    staff.get(f"/payments/{check}")
    downloaded(db, tmp_path, monkeypatch, check)
    r = staff.post(f"/payments/{check}/approve", data={"csrf": token, "transaction_id": "TX1", "amount": "18000",
                                                        "confirmed": "1"}, follow_redirects=False)
    assert r.status_code == 303 and order(db, "91")[0] == "paid"  # setting ON (default)
    deliver(db)

    owner.post("/users/staff-settings", data={"csrf": owner_token})  # switch it OFF
    check = _photo_check(owner, owner_token, db, "92", 2)
    page = staff.get(f"/payments/{check}").text
    assert "only the owner checks payments now" in page and f'action="/payments/{check}/approve"' not in page
    downloaded(db, tmp_path, monkeypatch, check, content=b"other")
    for action in ("approve", "reject", "not-payment", "ask-one"):
        r = staff.post(f"/payments/{check}/{action}", data={"csrf": token, "transaction_id": "TX2", "amount": "18000",
                                                              "reason": "not_received"}, follow_redirects=False)
        assert r.status_code == 403, action
    assert order(db, "92")[0] == "photo_received"
    r = owner.post(f"/payments/{check}/approve", data={"csrf": owner_token, "transaction_id": "TX2", "amount": "18000",
                                                        "confirmed": "1"}, follow_redirects=False)
    assert r.status_code == 303 and order(db, "92")[0] == "paid"


def test_d3_staff_taught_pattern_owner_decides(shop, db):
    owner, owner_token = shop
    y_shop(db)
    staff, token = staff_client(db)
    teach(staff, token, say(db, 1, "ysl edp 10 ml ရောယူပါမယ်"), p0=str(pid(db, "Y Eau de Parfum")), nick0="ysl edp")
    assert learned(db)[0][:3] == ("nickname", "ysl edp", "pending")
    assert scalar(db, "SELECT u.username FROM learned_patterns l JOIN users u ON u.id = l.taught_by") == "staff1"
    assert "📘 1 pattern to approve" not in staff.get("/").text  # approving is the owner's job
    assert "📘 1 pattern to approve" in owner.get("/").text and "taught by staff1" in owner.get("/learning").text
    p = scalar(db, "SELECT id FROM learned_patterns")
    assert staff.post(f"/learning/{p}/decide", data={"csrf": token, "status": "approved"}).status_code == 403
    owner.post(f"/learning/{p}/decide", data={"csrf": owner_token, "status": "approved"})
    assert learned(db)[0][2] == "approved"
    mid = say(db, 2, "ysl edp 5ml ယူမယ်", customer="92")
    assert scalar(db, "SELECT i.match FROM parse_items i JOIN parse_results r ON r.id = i.parse_result_id "
                      "WHERE r.message_id = :m AND r.is_current", m=mid) == "learned"
    owner.post(f"/learning/{p}/decide", data={"csrf": owner_token, "status": "rejected"})  # remove it again
    mid = say(db, 3, "ysl edp 5ml ယူမယ်", customer="93")
    assert scalar(db, "SELECT i.match FROM parse_items i JOIN parse_results r ON r.id = i.parse_result_id "
                      "WHERE r.message_id = :m AND r.is_current", m=mid) == "fuzzy"  # ignored by the parser


def test_d3b_staff_alias_in_correction_waits_for_owner(shop, db):
    staff, token = staff_client(db)
    mid = say(db, 1, "sauvage noir 5ml ယူမယ်")
    sauvage = scalar(db, "SELECT id FROM products WHERE name = 'Sauvage Eau de Parfum'")
    r = staff.post(f"/messages/{mid}/parse", data={"csrf": token, "intent": "buy", "product_id": str(sauvage),
                                                   "size_ml": "5", "alias": "sauvage noir"})
    assert "waits for the owner&#39;s approval" in r.text
    assert scalar(db, "SELECT count(*) FROM product_aliases WHERE normalized_alias = 'sauvage noir'") == 0
    assert learned(db) == [("nickname", "sauvage noir", "pending", "Sauvage Eau de Parfum", 0, 0)]


def test_d4_auto_approve_zero_means_owner_only(shop, db):
    owner, owner_token = shop
    y_shop(db)
    owner.post("/settings/learning", data={"csrf": owner_token, "auto_approve_after": "0", "greeting": "1",
                                           "auto_not_stocked": "1"})
    teach(owner, owner_token, say(db, 1, "ysl edp 10 ml ရောယူပါမယ်"), p0=str(pid(db, "Y Eau de Parfum")),
          nick0="ysl edp")
    deliver(db)
    for n in range(2, 7):  # five accepted uses
        tap(owner, owner_token, say(db, n, "ysl edp 5ml ယူမယ်", customer=str(90 + n)))
        deliver(db)
    assert learned(db)[0][2:5] == ("pending", "Y Eau de Parfum", 5)


def test_d5_disabled_user_cannot_log_in(shop, db):
    owner, owner_token = shop
    staff, token = staff_client(db)
    uid = scalar(db, "SELECT id FROM users WHERE username = 'staff1'")
    owner.post(f"/users/{uid}/disabled", data={"csrf": owner_token, "disabled": "1"})
    r = staff.get("/", follow_redirects=False)  # the open session ends at the next click
    assert r.status_code == 303 and r.headers["location"].startswith("/login")
    fresh = TestClient(staff.app, base_url="https://testserver")
    r = fresh.post("/login", data={"username": "staff1", "password": PASSWORD, "csrf": csrf(fresh), "next": "/"},
                   follow_redirects=False)
    assert r.status_code == 401 and "Wrong username or password" in r.text
    owner.post(f"/users/{uid}/disabled", data={"csrf": owner_token, "disabled": "0"})
    r = fresh.post("/login", data={"username": "staff1", "password": PASSWORD, "csrf": csrf(fresh), "next": "/"},
                   follow_redirects=False)
    assert r.status_code == 303
    me = scalar(db, "SELECT id FROM users WHERE username = 'owner'")
    assert "can&#39;t disable yourself" in owner.post(f"/users/{me}/disabled",
                                                  data={"csrf": owner_token, "disabled": "1"}).text


def test_users_page_add_reset_language_and_activity(shop, db):
    owner, token = shop
    r = owner.post("/users", data={"csrf": token, "username": "Mya", "password": "long enough pw",
                                   "role": "staff", "language": "my"})
    assert "Staff “mya” added" in r.text
    assert scalar(db, "SELECT role || '/' || language FROM users WHERE username = 'mya'") == "staff/my"
    assert "at least 10" in owner.post("/users", data={"csrf": token, "username": "x", "password": "short"}).text
    uid = scalar(db, "SELECT id FROM users WHERE username = 'mya'")
    owner.post(f"/users/{uid}/password", data={"csrf": token, "password": "another long pw"})
    owner.post(f"/users/{uid}/language", data={"csrf": token, "language": "en"})
    assert scalar(db, "SELECT language FROM users WHERE id = :u", u=uid) == "en"
    mya = TestClient(owner.app, base_url="https://testserver")
    r = mya.post("/login", data={"username": "mya", "password": "another long pw", "csrf": csrf(mya), "next": "/"},
                 follow_redirects=False)
    assert r.status_code == 303
    # every action shows who did it
    mid = say(db, 1, "ကျေးဇူးပါ")
    mya.post(f"/messages/{mid}/review", data={"csrf": csrf(mya, "/"), "status": "done"})
    page = owner.get("/activity").text
    rows = re.findall(r"<tr><td class=\"nowrap\">[^<]*</td><td>([^<]+)</td>\s*<td>([^<]+)</td>", page)
    assert ("mya", "message.review") in rows and ("owner", "user.add") in rows
    assert "mya" in owner.get(f"/activity?who={uid}").text


def test_create_user_script_role(monkeypatch, db):
    from scripts import create_user

    answers = iter(["staff password 1", "staff password 1"])
    monkeypatch.setattr(create_user.getpass, "getpass", lambda prompt: next(answers))
    monkeypatch.setattr("sys.argv", ["create_user", "Thida", "--role", "staff"])
    create_user.main()
    with SessionLocal() as s:
        u = s.query(User).filter_by(username="thida").one()
        assert (u.role, u.language) == ("staff", "my")
