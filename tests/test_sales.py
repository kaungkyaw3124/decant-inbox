"""docs/SPEC-v3.md Step 8 acceptance tests C1-C5: sales & profit (paid orders only, Yangon
days), cost per ml + vial cost, price at the time of sale, missing cost, export = page."""
import io
import re
from datetime import UTC, date, datetime

from openpyxl import load_workbook
from sqlalchemy import text

from app.core.db import SessionLocal
from app.orders import sales
from app.pricing import excel
from tests.test_after_payment import paid_order
from tests.test_cart import shop  # noqa: F401  (fixture)
from tests.test_messaging import scalar
from tests.test_payments import waiting_for_payment
from tests.test_prices import HEADER, STW, run_import, xlsx
from tests.test_web import client, log_in  # noqa: F401  (fixture)

WITH_COST = HEADER + ["bottle_cost"]


def set_paid_at(db, order_id: int, when: datetime) -> None:
    with db.begin() as c:
        c.execute(text("UPDATE orders SET paid_at = :w WHERE id = :o"), {"w": when, "o": order_id})


def report(first: date, last: date) -> sales.Report:
    with SessionLocal() as s:
        return sales.report(s, "custom", first, last)


def costs(db, client, token, bottle_cost=300000, vial_5ml="500") -> None:
    run_import(xlsx(STW + [bottle_cost], header=WITH_COST))
    client.post("/settings/vial-costs", data={"csrf": token, "vial_5": vial_5ml})


def test_c1_only_paid_orders_by_yangon_day(shop, db, tmp_path, monkeypatch):
    client, token = shop
    a = paid_order(client, token, db, tmp_path, monkeypatch, "91", 1)
    b = paid_order(client, token, db, tmp_path, monkeypatch, "92", 2)
    c = paid_order(client, token, db, tmp_path, monkeypatch, "93", 3)
    client.post(f"/orders/{c}/cancel", data={"csrf": token, "note": "refunded", "message": "none"})
    waiting_for_payment(client, token, db, customer="94")  # not paid
    # Yangon = UTC+6:30: 17:29 UTC is 23:59 on the 14th there, 17:31 UTC is 00:01 on the 15th
    set_paid_at(db, a, datetime(2026, 10, 14, 17, 29, tzinfo=UTC))
    set_paid_at(db, b, datetime(2026, 10, 14, 17, 31, tzinfo=UTC))
    set_paid_at(db, c, datetime(2026, 10, 14, 18, 0, tzinfo=UTC))
    r = report(date(2026, 10, 14), date(2026, 10, 15))
    assert (r.total.revenue, len(r.total.orders), r.total.items, r.total.ml) == (36000, 2, 2, 10)
    assert [(d.isoformat(), g.revenue, len(g.orders)) for d, g in r.days] == [
        ("2026-10-14", 18000, 1), ("2026-10-15", 18000, 1)]
    assert [(g.label, g.items, g.revenue) for g in r.perfumes] == [("Stronger With You Intensely", 2, 36000)]
    assert [(g.label, g.items) for g in r.sizes] == [("5ml", 2)]
    assert report(date(2026, 10, 15), date(2026, 10, 15)).total.revenue == 18000

    page = client.get("/sales?range=custom&from=2026-10-14&to=2026-10-15").text
    assert "36,000" in page and "Revenue per day" in page and "<svg" in page
    assert "Sun 15 Oct: 18,000 ကျပ် · 1 order(s)" not in page  # 15 Oct 2026 is a Thursday
    assert "Thu 15 Oct: 18,000 ကျပ် · 1 order(s)" in page
    assert client.get("/sales").status_code == 200 and client.get("/sales?range=month").status_code == 200


def test_c2_cost_per_ml_and_vial(shop, db, tmp_path, monkeypatch):
    client, token = shop
    costs(db, client, token)  # bottle 300,000 for 100 ml, 5 ml vial 500
    assert scalar(db, "SELECT bottle_cost_mmk FROM product_costs") == 300000
    order_id = paid_order(client, token, db, tmp_path, monkeypatch, "91", 1)
    assert scalar(db, "SELECT unit_cost_mmk FROM order_items WHERE order_id = :o", o=order_id) == 15500
    today = sales.today()
    t = report(today, today).total
    assert (t.revenue, t.cost, t.profit, round(t.margin, 1), t.average) == (18000, 15500, 2500, 13.9, 18000)
    page = client.get("/sales").text
    assert "15,500" in page and "2,500" in page and "13.9%" in page


def test_c3_report_uses_price_and_cost_at_sale(shop, db, tmp_path, monkeypatch):
    client, token = shop
    costs(db, client, token)
    paid_order(client, token, db, tmp_path, monkeypatch, "91", 1)
    changed = list(STW)
    changed[5] = 21000  # 5ml price up after the sale
    run_import(xlsx(changed + [360000], header=WITH_COST))  # and the bottle got dearer
    today = sales.today()
    t = report(today, today).total
    assert (t.revenue, t.cost, t.profit) == (18000, 15500, 2500)


def test_c4_missing_cost_shows_dash(shop, db, tmp_path, monkeypatch):
    client, token = shop
    paid_order(client, token, db, tmp_path, monkeypatch, "91", 1)  # no bottle_cost yet
    today = sales.today()
    t = report(today, today).total
    assert t.revenue == 18000 and t.cost is None and t.profit is None and t.margin is None
    page = client.get("/sales").text
    assert "<b>—</b>" in page and "add bottle_cost in the" in page
    # fixing it: add the cost, then use today's cost for the order already paid
    costs(db, client, token)
    client.post("/sales/fill-costs", data={"csrf": token, "range": "today"})
    assert report(today, today).total.profit == 2500


def test_c5_export_matches_page(shop, db, tmp_path, monkeypatch):
    client, token = shop
    costs(db, client, token)
    paid_order(client, token, db, tmp_path, monkeypatch, "91", 1)
    paid_order(client, token, db, tmp_path, monkeypatch, "92", 2)
    page = client.get("/sales?range=week").text
    tiles = dict(re.findall(r'<div class="tile"><span>([^<]+)</span><b>([^<]+)</b>', page))
    r = client.get("/sales/export.xlsx?range=week")
    assert r.headers["content-disposition"].startswith('attachment; filename="sales-')
    wb = load_workbook(io.BytesIO(r.content))
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True)}
    assert f"{summary['Revenue (ကျပ်)']:,}" == tiles["Revenue"] == "36,000"
    assert f"{summary['Cost (ကျပ်)']:,}" == tiles["Cost"] == "31,000"
    assert f"{summary['Profit (ကျပ်)']:,}" == tiles["Profit"] == "5,000"
    assert str(summary["Orders"]) == tiles["Orders"] and str(summary["Items sold"]) == tiles["Items sold"]
    assert f"{summary['Profit %']}%" == tiles["Profit %"]
    perfume = list(wb["Per perfume"].iter_rows(values_only=True))
    assert perfume[1] == ("Stronger With You Intensely", 2, 2, 10, 36000, 31000, 5000, 13.9)
    assert perfume[-1][0] == "Total" and perfume[-1][4] == 36000
    days = list(wb["Per day"].iter_rows(values_only=True))
    assert sum(row[4] for row in days[1:-1]) == 36000


def test_excel_bottle_cost_roundtrip(shop, db):
    client, token = shop
    run_import(xlsx(STW + [300000], header=WITH_COST))
    with SessionLocal() as s:
        rows, sizes = excel.export_rows(s)
    sheet = excel.read_sheet(excel.write_workbook(rows, sizes))
    assert {r.name: r.bottle_cost for r in sheet.rows}["Stronger With You Intensely"] == 300000
    with SessionLocal() as s:
        plan = excel.plan_import(s, sheet)
    assert all(r.status == "same" for r in plan.rows)
    files = {"file": ("p.xlsx", xlsx(STW + [320000], header=WITH_COST), "application/octet-stream")}
    assert "bottle cost: 300,000 → 320,000" in client.post("/prices/import", data={"csrf": token}, files=files).text
    bad = excel.read_sheet(xlsx(STW + ["lots"], header=WITH_COST))
    assert any("is not a cost" in e for e in bad.errors)


def test_unit_cost_math():
    from app.pricing.costs import unit_cost

    assert unit_cost(100, 5, False, 300000, {5: 500}) == 15500
    assert unit_cost(100, None, True, 300000, {5: 500}) == 300000  # full bottle: no vial
    assert unit_cost(None, 5, False, 300000, {}) is None and unit_cost(100, 5, False, None, {}) is None
