"""Daily sales & profit (docs/SPEC-v3.md Step 8).

Counts only paid orders (paid, packed, sent out, delivered) by the day they were paid
(Asia/Yangon days). Revenue = the prices on the order items (the price at the time of the
sale); cost = the cost stamped on the items when the order was paid. If any item in a
group has no cost, that group's cost and profit are unknown ("—").
"""
import io
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time, timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import Order, OrderItem, Product
from app.models.orders import PAID_STATUSES
from app.pricing.prices import size_label

RANGES = {"today": "Today", "yesterday": "Yesterday", "week": "This week", "month": "This month",
          "custom": "Custom"}
MAX_DAYS = 366
MISSING = "—"


def tz() -> ZoneInfo:
    return ZoneInfo(get_settings().timezone)


def today() -> date:
    return datetime.now(tz()).date()


def resolve(key: str, start: str = "", end: str = "", now_date: date | None = None) -> tuple[str, date, date]:
    """(range key, first day, last day). Unknown / broken input falls back to today."""
    d = now_date or today()
    if key == "yesterday":
        return key, d - timedelta(days=1), d - timedelta(days=1)
    if key == "week":  # Monday to today
        return key, d - timedelta(days=d.weekday()), d
    if key == "month":
        return key, d.replace(day=1), d
    if key == "custom":
        try:
            first, last = date.fromisoformat(start), date.fromisoformat(end)
        except ValueError:
            return "today", d, d
        first, last = min(first, last), max(first, last)
        return key, max(first, last - timedelta(days=MAX_DAYS - 1)), last
    return "today", d, d


def bounds(first: date, last: date) -> tuple[datetime, datetime]:
    """UTC [start, end) of Yangon days first..last."""
    start = datetime.combine(first, time(0), tzinfo=tz()).astimezone(UTC)
    end = datetime.combine(last + timedelta(days=1), time(0), tzinfo=tz()).astimezone(UTC)
    return start, end


@dataclass
class Group:
    label: str
    orders: set[int] = field(default_factory=set)
    items: int = 0
    ml: int = 0
    revenue: int = 0
    known_cost: int = 0
    missing_cost: bool = False

    def add(self, order_id: int, qty: int, ml: int, revenue: int, cost: int | None) -> None:
        self.orders.add(order_id)
        self.items += qty
        self.ml += ml
        self.revenue += revenue
        if cost is None:
            self.missing_cost = True
        else:
            self.known_cost += cost

    @property
    def cost(self) -> int | None:
        return None if self.missing_cost else self.known_cost

    @property
    def profit(self) -> int | None:
        return None if self.missing_cost else self.revenue - self.known_cost

    @property
    def margin(self) -> float | None:
        profit = self.profit
        return None if profit is None or not self.revenue else profit * 100 / self.revenue

    @property
    def average(self) -> int:
        return round(self.revenue / len(self.orders)) if self.orders else 0


@dataclass
class Report:
    key: str
    first: date
    last: date
    total: Group
    days: list[tuple[date, Group]]
    perfumes: list[Group]
    sizes: list[Group]
    missing_orders: int  # paid orders in the range with an item without cost

    @property
    def label(self) -> str:
        if self.first == self.last:
            return self.first.strftime("%d %b %Y")
        return f"{self.first:%d %b} – {self.last:%d %b %Y}"


def report(session: Session, key: str, first: date, last: date) -> Report:
    start, end = bounds(first, last)
    rows = session.execute(
        select(OrderItem, Order, Product).join(Order, Order.id == OrderItem.order_id)
        .join(Product, Product.id == OrderItem.product_id)
        .where(Order.status.in_(PAID_STATUSES), Order.paid_at >= start, Order.paid_at < end,
               OrderItem.removed_at.is_(None))
        .order_by(Order.paid_at, OrderItem.id)).all()
    total = Group("Total")
    days = {first + timedelta(days=n): Group("") for n in range((last - first).days + 1)}
    perfumes: dict[int, Group] = {}
    sizes: dict[tuple, Group] = {}
    missing: set[int] = set()
    for item, order, product in rows:
        qty = item.quantity
        ml = (product.full_size_ml or 0 if item.is_full_bottle else item.size_ml or 0) * qty
        revenue = (item.unit_price_mmk or 0) * qty
        cost = item.unit_cost_mmk * qty if item.unit_cost_mmk is not None else None
        if cost is None:
            missing.add(order.id)
        day = order.paid_at.astimezone(tz()).date()
        size_key = (1, 0) if item.is_full_bottle else (0, item.size_ml or 0)
        for group in (total, days[day], perfumes.setdefault(product.id, Group(product.name)),
                      sizes.setdefault(size_key, Group("Full bottle" if item.is_full_bottle
                                                       else size_label(item.size_ml, None)))):
            group.add(order.id, qty, ml, revenue, cost)
    return Report(key, first, last, total, sorted(days.items()),
                  sorted(perfumes.values(), key=lambda g: (-g.revenue, g.label)),
                  [g for _, g in sorted(sizes.items())], len(missing))


def unstamped_orders(session: Session, first: date, last: date) -> list[Order]:
    start, end = bounds(first, last)
    return list(session.scalars(select(Order).join(OrderItem, OrderItem.order_id == Order.id).where(
        Order.status.in_(PAID_STATUSES), Order.paid_at >= start, Order.paid_at < end,
        OrderItem.removed_at.is_(None), OrderItem.unit_cost_mmk.is_(None)).distinct()))


# ---------- chart (inline SVG, no libraries) ----------

CHART_W, CHART_H, CHART_TOP, CHART_BOTTOM = 600, 170, 14, 22


def chart(r: Report) -> dict:
    """Bars of revenue per day, positioned for the template."""
    n = len(r.days)
    peak = max((g.revenue for _, g in r.days), default=0)
    slot = CHART_W / max(n, 1)
    width = max(min(slot * 0.7, 48), 2)
    every = max(1, round(n / 12))  # at most ~12 day labels
    plot_h = CHART_H - CHART_TOP - CHART_BOTTOM
    bars = []
    for i, (day, g) in enumerate(r.days):
        h = plot_h * g.revenue / peak if peak else 0
        if g.revenue and h < 2:
            h = 2
        x = slot * i + (slot - width) / 2
        bars.append({"x": round(x, 1), "y": round(CHART_TOP + plot_h - h, 1), "w": round(width, 1),
                     "h": round(h, 1), "cx": round(x + width / 2, 1),
                     "hx": round(slot * i, 1), "hw": round(slot, 1),  # hover target = the whole slot
                     "label": f"{day:%d}" if i % every == 0 else "",
                     "title": f"{day:%a %d %b}: {g.revenue:,} ကျပ် · {len(g.orders)} order(s)"})
    return {"w": CHART_W, "h": CHART_H, "base": CHART_TOP + plot_h, "bars": bars, "peak": peak}


# ---------- export ----------

def _money(value: int | None):
    return MISSING if value is None else value


def _pct(value: float | None):
    return MISSING if value is None else round(value, 1)


def export(r: Report) -> bytes:
    """The same numbers as the page, one sheet per table."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    t = r.total
    for row in (("Sales", r.label), ("From", r.first.isoformat()), ("To", r.last.isoformat()),
                ("Revenue (ကျပ်)", t.revenue), ("Orders", len(t.orders)), ("Items sold", t.items),
                ("ml sold", t.ml), ("Cost (ကျပ်)", _money(t.cost)), ("Profit (ကျပ်)", _money(t.profit)),
                ("Profit %", _pct(t.margin)), ("Average order (ကျပ်)", t.average)):
        ws.append(row)
    if t.missing_cost:
        ws.append(("Note", "Some items have no cost: add bottle_cost in the price list."))
    header = ["Orders", "Items", "ml", "Revenue", "Cost", "Profit", "Profit %"]

    def values(g: Group) -> list:
        return [len(g.orders), g.items, g.ml, g.revenue, _money(g.cost), _money(g.profit), _pct(g.margin)]

    for title, first_col, groups in (
            ("Per day", "Day", [(d.isoformat(), g) for d, g in r.days]),
            ("Per perfume", "Perfume", [(g.label, g) for g in r.perfumes]),
            ("Per size", "Size", [(g.label, g) for g in r.sizes])):
        sheet = wb.create_sheet(title)
        sheet.append([first_col, *header])
        for label, g in groups:
            sheet.append([label, *values(g)])
        sheet.append(["Total", *values(t)])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.column_dimensions["A"].width = 34
        for col in range(2, len(header) + 2):
            sheet.column_dimensions[get_column_letter(col)].width = 13
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 26
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
