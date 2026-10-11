"""Price list as Excel: one row per perfume.

    brand | name | full_size_ml | full_price | 2ml | 5ml | 10ml | ... | in_stock

- Any header like "3ml" / "15 ml" is a decant size. Empty cell = size not sold.
- A size column that is missing from the file leaves that size unchanged.
- in_stock: yes/no; empty = unchanged (yes for new products). Applies to every size.
- Import is two steps: read_sheet() + plan_import() build a preview without writing
  anything; apply_import() writes it after the admin confirms.
"""
import io
import re
from dataclasses import asdict, dataclass, field
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.models import Product, ProductAlias, ProductPrice
from app.parser.normalize import normalize
from app.pricing.prices import DEFAULT_SIZES, format_mmk, latest_prices, size_label

SHEET_NAME = "Prices"
MAX_ROWS = 2000
MAX_COLUMNS = 40
MAX_PRICE = 100_000_000
MAX_SIZE_ML = 1000
FUZZY_THRESHOLD = 80
NAMED_COLUMNS = ("brand", "name", "full_size_ml", "full", "in_stock")  # "full" = full_price
YES = {"yes", "y", "true", "1", "ရှိ", "ရှိတယ်"}
NO = {"no", "n", "false", "0", "မရှိ", "ကုန်", "ကုန်ပြီ"}
_SIZE_HEADER = re.compile(r"^(\d+)\s*ml$")
_MONEY_NOISE = re.compile(r"(,|\s|ks|kyats?|mmk|ကျပ်|ml)", re.IGNORECASE)

HELP_LINES = [
    "How to use this price list",
    "",
    "One row per perfume. Prices in kyat, numbers only (18000 or 18,000).",
    "Size columns: any header like 2ml, 5ml, 10ml. Add a column (e.g. 3ml or 15ml) to sell a new size.",
    "Empty price cell = that size is not sold. Deleting a whole size column leaves that size unchanged.",
    "full_size_ml = bottle size (e.g. 100), full_price = price of the full bottle.",
    "in_stock: yes or no (for the whole perfume). Empty = unchanged.",
    "A new brand + name adds a new product. Names that look like an existing product are shown for you to decide.",
    "Upload in the dashboard: Prices > Upload. You see a preview first; nothing is saved until you press Confirm.",
]


class SheetError(Exception):
    """The file can't be read at all (not xlsx, no header, too big)."""


@dataclass
class SheetRow:
    row: int  # Excel row number, for messages
    brand: str
    name: str
    full_size_ml: int | None  # only meaningful if the file has a full_size_ml column
    prices: dict[str, int | None]  # "5" -> 18000, "full" -> 280000; only columns in the file
    in_stock: bool | None  # None = unchanged
    full_size_given: bool = False


@dataclass
class Sheet:
    rows: list[SheetRow] = field(default_factory=list)
    has_full_size: bool = False
    errors: list[str] = field(default_factory=list)  # block the import
    warnings: list[str] = field(default_factory=list)

    def to_payload(self) -> dict[str, Any]:
        return {"rows": [asdict(r) for r in self.rows], "has_full_size": self.has_full_size}

    @classmethod
    def from_payload(cls, data: dict[str, Any]) -> "Sheet":
        return cls(rows=[SheetRow(**r) for r in data["rows"]], has_full_size=data["has_full_size"])


def size_key(key: str) -> int | None:
    return None if key == "full" else int(key)


# ---------- reading ----------

def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _number(value: Any, maximum: int) -> int | None:
    """Positive whole number from a cell. None for empty / "-". ValueError if invalid."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError
        number = int(value)
    else:
        text = _MONEY_NOISE.sub("", normalize(str(value)))
        if text in ("", "-"):
            return None
        if not text.isdigit():
            raise ValueError
        number = int(text)
    if not 0 < number <= maximum:
        raise ValueError
    return number


def _stock(value: Any) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    text = normalize(_text(value))
    if text in YES:
        return True
    if text in NO:
        return False
    raise ValueError


def _header_key(value: Any) -> str:
    text = normalize(_text(value))
    if m := _SIZE_HEADER.match(text):
        return m.group(1)
    if text in ("full", "full price", "full_price"):
        return "full"
    return text.replace(" ", "_")


def _pick_sheet(workbook):
    return workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]


def read_sheet(data: bytes) -> Sheet:
    try:
        values_wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas_wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises many types for broken files
        raise SheetError("This is not a valid .xlsx file. Save it as Excel Workbook (.xlsx) and try again.") from exc
    try:
        values = list(_pick_sheet(values_wb).iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1,
                                                       values_only=True))
        formulas = list(_pick_sheet(formulas_wb).iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1,
                                                           values_only=True))
    finally:
        values_wb.close()
        formulas_wb.close()
    if not values:
        raise SheetError("The sheet is empty.")
    if len(values) > MAX_ROWS + 1:
        raise SheetError(f"Too many rows (max {MAX_ROWS} perfumes).")

    sheet = Sheet()
    header = [_header_key(v) for v in values[0]]
    while header and header[-1] == "":
        header.pop()
    if len(header) > MAX_COLUMNS:
        raise SheetError(f"Too many columns (max {MAX_COLUMNS}).")
    columns: dict[str, int] = {}
    for idx, key in enumerate(header):
        letter = get_column_letter(idx + 1)
        if key == "":
            continue
        if key in columns:
            sheet.errors.append(f"Column {letter}: '{_text(values[0][idx])}' appears twice.")
        elif key.isdigit() or key in NAMED_COLUMNS:
            columns[key] = idx
        else:
            sheet.warnings.append(f"Column {letter} '{_text(values[0][idx])}' is ignored "
                                  "(not brand/name/full_size_ml/full_price/in_stock or a size like 5ml).")
    for required in ("brand", "name"):
        if required not in columns:
            raise SheetError(f"Header row (row 1) must have a '{required}' column.")
    for key in columns:
        if key.isdigit() and not 0 < int(key) <= MAX_SIZE_ML:
            sheet.errors.append(f"Size column '{key}ml' is not a valid size.")
    sheet.has_full_size = "full_size_ml" in columns
    price_keys = sorted((k for k in columns if k.isdigit()), key=int) + (["full"] if "full" in columns else [])

    seen: dict[tuple[str, str], int] = {}
    for offset, (row_values, row_formulas) in enumerate(zip(values[1:], formulas[1:]), start=2):
        def cell(key: str) -> Any:
            idx = columns.get(key)
            return row_values[idx] if idx is not None and idx < len(row_values) else None

        def ref(key: str) -> str:
            return f"{get_column_letter(columns[key] + 1)}{offset}"

        if all(_text(v) == "" for v in row_values):
            continue
        for key, idx in columns.items():
            if (idx < len(row_formulas) and isinstance(row_formulas[idx], str)
                    and row_formulas[idx].startswith("=") and row_values[idx] is None):
                sheet.errors.append(f"{ref(key)}: formula without a saved result. Open the file in Excel, "
                                    "save it, and upload again (or type the number).")
        brand, name = _text(cell("brand")), _text(cell("name"))
        if not brand or not name:
            sheet.errors.append(f"Row {offset}: brand and name are both required.")
            continue
        dup_key = (normalize(brand), normalize(name))
        if dup_key in seen:
            sheet.errors.append(f"Row {offset}: {brand} {name} is already on row {seen[dup_key]}.")
            continue
        seen[dup_key] = offset

        prices: dict[str, int | None] = {}
        for key in price_keys:
            try:
                prices[key] = _number(cell(key), MAX_PRICE)
            except ValueError:
                sheet.errors.append(f"{ref(key)}: '{_text(cell(key))}' is not a price (whole kyat, e.g. 18000).")
        full_size = None
        if sheet.has_full_size:
            try:
                full_size = _number(cell("full_size_ml"), MAX_SIZE_ML)
            except ValueError:
                sheet.errors.append(f"{ref('full_size_ml')}: '{_text(cell('full_size_ml'))}' is not a bottle "
                                    "size in ml (e.g. 100).")
        in_stock = None
        if "in_stock" in columns:
            try:
                in_stock = _stock(cell("in_stock"))
            except ValueError:
                sheet.errors.append(f"{ref('in_stock')}: '{_text(cell('in_stock'))}' must be yes or no.")
        if prices.get("full") and sheet.has_full_size and not full_size:
            sheet.warnings.append(f"Row {offset}: {name} has a full bottle price but no full_size_ml.")
        sheet.rows.append(SheetRow(offset, brand, name, full_size, prices, in_stock, sheet.has_full_size))
    if not sheet.rows and not sheet.errors:
        sheet.errors.append("No perfumes found below the header row.")
    return sheet


# ---------- preview ----------

@dataclass
class RowPlan:
    row: SheetRow
    status: str  # changed | same | new | ambiguous
    product_id: int | None = None
    candidates: list[tuple[int, str]] = field(default_factory=list)  # (product_id, "Brand Name")
    changes: list[str] = field(default_factory=list)  # human-readable lines
    reactivate: bool = False


@dataclass
class Plan:
    rows: list[RowPlan]
    untouched: list[str]  # active products that are not in the file

    def count(self, status: str) -> int:
        return sum(1 for r in self.rows if r.status == status)


def _money(value: int | None) -> str:
    return "not sold" if value is None else format_mmk(value)


def _desired(row: SheetRow, current: dict[int | None, Any]) -> dict[int | None, tuple[int | None, bool]]:
    """Final (price, in_stock) per size after the import, for sizes that would change."""
    final: dict[int | None, tuple[int | None, bool]] = {
        size: (p.price_mmk, p.in_stock) for size, p in current.items() if p.price_mmk is not None}
    # a newly sold size gets the perfume's stock: out only if every size is out
    product_in_stock = not final or any(stock for _, stock in final.values())
    for key, price in row.prices.items():
        size = size_key(key)
        if price is None:
            final.pop(size, None)
        else:
            final[size] = (price, final.get(size, (None, product_in_stock))[1])
    if row.in_stock is not None:
        final = {size: (price, row.in_stock) for size, (price, _) in final.items()}
    result = {size: v for size, v in final.items()
              if size not in current or (current[size].price_mmk, current[size].in_stock) != v}
    for size, p in current.items():  # sizes removed by an empty cell
        if p.price_mmk is not None and size not in final:
            result[size] = (None, p.in_stock)
    return result


def _bottle_size_changes(row: SheetRow, product: Product) -> bool:
    return row.full_size_given and row.full_size_ml != product.full_size_ml


def _describe(row: SheetRow, product: Product | None, current: dict[int | None, Any]) -> list[str]:
    lines = []
    full_size = row.full_size_ml if row.full_size_given or product is None else product.full_size_ml
    if product is not None and _bottle_size_changes(row, product):
        lines.append(f"bottle size: {product.full_size_ml or '?'}ml → {row.full_size_ml or '?'}ml")
    changes = _desired(row, current)
    old_stock = {p.in_stock for p in current.values() if p.price_mmk is not None}
    for size in sorted(changes, key=lambda s: (s is None, s or 0)):
        new_price, new_stock = changes[size]
        old = current.get(size)
        old_price = old.price_mmk if old else None
        if old_price != new_price:
            label = size_label(size, full_size)
            lines.append(f"{label}: {_money(new_price)}" if old_price is None and product is None
                         else f"{label}: {_money(old_price)} → {_money(new_price)}")
    if row.in_stock is not None and old_stock and old_stock != {row.in_stock}:
        lines.append(f"in stock: {'yes' if True in old_stock else 'no'} → {'yes' if row.in_stock else 'no'}")
    elif row.in_stock is False and not old_stock:
        lines.append("in stock: no")
    return lines


def _candidates(row: SheetRow, products: list[Product]) -> list[tuple[int, str]]:
    target = normalize(f"{row.brand} {row.name}")
    name = normalize(row.name)
    scored = []
    for p in products:
        score = max(fuzz.token_sort_ratio(target, normalize(f"{p.brand} {p.name}")),
                    fuzz.ratio(name, normalize(p.name)))
        if score >= FUZZY_THRESHOLD:
            scored.append((score, p.id, f"{p.brand} {p.name}"))
    return [(pid, label) for _, pid, label in sorted(scored, key=lambda s: (-s[0], s[2]))[:3]]


def plan_import(session: Session, sheet: Sheet) -> Plan:
    products = list(session.scalars(select(Product)))
    by_key = {(normalize(p.brand), normalize(p.name)): p for p in products}
    latest = latest_prices(session)
    active = [p for p in products if p.active]
    rows, matched = [], set()
    for row in sheet.rows:
        product = by_key.get((normalize(row.brand), normalize(row.name)))
        if product is None:
            candidates = _candidates(row, active)
            rows.append(RowPlan(row, "ambiguous" if candidates else "new", candidates=candidates,
                                changes=_describe(row, None, {})))
            continue
        matched.add(product.id)
        changes = _describe(row, product, latest.get(product.id, {}))
        rows.append(RowPlan(row, "changed" if changes or not product.active else "same", product.id,
                            changes=changes, reactivate=not product.active))
    untouched = sorted(f"{p.brand} {p.name}" for p in active if p.id not in matched)
    return Plan(rows, untouched)


def fingerprint(plan: Plan) -> list:
    """What the admin saw in the preview; confirm refuses if it no longer matches."""
    return [[r.row.row, r.status, r.product_id, r.changes] for r in plan.rows]


# ---------- apply ----------

def apply_import(session: Session, plan: Plan, choices: dict[int, str], user_id: int | None) -> dict[str, int]:
    """choices: Excel row number -> "new" | "skip" | "<product_id>" for new/ambiguous rows.
    Raises ValueError (nothing written) if a choice is missing or two rows hit one product."""
    targets: list[tuple[RowPlan, Product | None]] = []
    used: dict[int, int] = {r.product_id: r.row.row for r in plan.rows if r.product_id}
    for rp in plan.rows:
        if rp.status not in ("new", "ambiguous"):
            targets.append((rp, session.get(Product, rp.product_id)))
            continue
        choice = choices.get(rp.row.row, "")
        if choice == "skip":
            continue
        if choice == "new":
            targets.append((rp, None))
        elif choice.isdigit() and int(choice) in {pid for pid, _ in rp.candidates}:
            pid = int(choice)
            if pid in used:
                raise ValueError(f"Row {rp.row.row} and row {used[pid]} are both the same product.")
            used[pid] = rp.row.row
            targets.append((rp, session.get(Product, pid)))
        else:
            raise ValueError(f"Row {rp.row.row} ({rp.row.brand} {rp.row.name}): choose what to do with it.")

    now = datetime.now(UTC)
    latest = latest_prices(session)
    stats = {"products_added": 0, "products_updated": 0, "prices_written": 0}
    for rp, product in targets:
        row = rp.row
        if product is None:
            product = Product(brand=row.brand, name=row.name, source="excel", active=True,
                              full_size_ml=row.full_size_ml)
            session.add(product)
            session.flush()
            session.execute(insert(ProductAlias).values(
                product_id=product.id, alias_text=row.name, normalized_alias=normalize(row.name),
                source="learned").on_conflict_do_nothing())
            stats["products_added"] += 1
            added = True
        else:
            added = False
        changed = not product.active
        product.active = True
        if _bottle_size_changes(row, product):
            product.full_size_ml = row.full_size_ml
            changed = True
        for size, (price, in_stock) in _desired(row, latest.get(product.id, {})).items():
            session.add(ProductPrice(product_id=product.id, size_ml=size, is_full_bottle=size is None,
                                     price_mmk=price, in_stock=in_stock, effective_from=now, created_by=user_id))
            stats["prices_written"] += 1
            changed = True
        if changed and not added:
            stats["products_updated"] += 1
    return stats


# ---------- writing ----------

def write_workbook(rows: list[dict[str, Any]], sizes: list[int]) -> bytes:
    """rows: brand, name, full_size_ml, full_price, prices {size_ml: price}, in_stock (bool)."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    header = ["brand", "name", "full_size_ml", "full_price", *[f"{s}ml" for s in sizes], "in_stock"]
    ws.append(header)
    for r in rows:
        ws.append([r["brand"], r["name"], r.get("full_size_ml"), r.get("full_price"),
                   *[r["prices"].get(s) for s in sizes], "yes" if r.get("in_stock", True) else "no"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in range(4, 5 + len(sizes)):
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = "#,##0"
    widths = [22, 34, 13, 13, *[10] * len(sizes), 10]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "C2"
    help_ws = wb.create_sheet("Help")
    for line in HELP_LINES:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True)
    help_ws.column_dimensions["A"].width = 110
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_rows(session: Session) -> tuple[list[dict[str, Any]], list[int]]:
    products = list(session.scalars(select(Product).where(Product.active).order_by(Product.brand, Product.name)))
    latest = latest_prices(session, [p.id for p in products])
    sizes = set(DEFAULT_SIZES)
    rows = []
    for p in products:
        current = {s: pr for s, pr in latest.get(p.id, {}).items() if pr.price_mmk is not None}
        sizes.update(s for s in current if s is not None)
        full = current.get(None)
        rows.append({
            "brand": p.brand, "name": p.name, "full_size_ml": p.full_size_ml,
            "full_price": full.price_mmk if full else None,
            "prices": {s: pr.price_mmk for s, pr in current.items() if s is not None},
            # one stock flag per perfume in the file: "no" only if every size is out of stock
            "in_stock": not current or any(pr.in_stock for pr in current.values()),
        })
    return rows, sorted(sizes)
